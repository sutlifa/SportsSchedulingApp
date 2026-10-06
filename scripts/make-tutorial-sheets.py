#!/usr/bin/env python3
"""Writes the guide tutorial's sample facility spreadsheets, one per sport:
public/tutorial/<sport>-april-2027.xlsx and .csv.

    python3 scripts/make-tutorial-sheets.py      (needs: pip install openpyxl)

Each sheet lists the tutorial's first facility's April availability one row
per playing area, named the sport's way (Sheet A / Court 1 / Field 1), with a
"Tournament" and a "Reserved" closure. The tutorial promises exactly
"12 ... slots on 5 dates"; scripts/e2e/tutorial.mjs checks it. Keep these
names and numbers in step with lib/engine/sports.ts and the guide.
"""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font

SPORTS = {
    # id: (facility, unit heading, unit labels, booked-time word)
    "tennis": ("Riverside Tennis Center", "Court", ["Court 1", "Court 2", "Court 3"], "court time"),
    "pickleball": ("Riverside Pickleball Club", "Court", ["Court 1", "Court 2", "Court 3"], "court time"),
    "hockey": ("Riverside Ice Arena", "Sheet", ["Sheet A", "Sheet B", "Sheet C"], "ice time"),
    "soccer": ("Riverside Soccer Complex", "Field", ["Field 1", "Field 2", "Field 3"], "field time"),
    "basketball": ("Riverside Community Gym", "Court", ["Court 1", "Court 2", "Court 3"], "gym time"),
    "volleyball": ("Riverside Sports Center", "Court", ["Court 1", "Court 2", "Court 3"], "court time"),
    "baseball": ("Riverside Ballpark", "Field", ["Field 1", "Field 2", "Field 3"], "field time"),
    "softball": ("Riverside Softball Complex", "Field", ["Field 1", "Field 2", "Field 3"], "field time"),
    "lacrosse": ("Riverside Sports Complex", "Field", ["Field 1", "Field 2", "Field 3"], "field time"),
    "football": ("Riverside Stadium", "Field", ["Field 1", "Field 2", "Field 3"], "field time"),
    "other": ("Riverside Sports Center", "Space", ["Space 1", "Space 2", "Space 3"], "facility time"),
}

# (date, start, unit index or a closure word). 12 time slots on 5 dates.
PLAN = [
    ("Sat 4/3", "9:00 AM", [0, 1, 2]), ("Sat 4/3", "11:00 AM", [0, 1]), ("Sat 4/3", "1:00 PM", [0, 1, 2]),
    ("Sat 4/10", "9:00 AM", [0, 1]), ("Sat 4/10", "11:00 AM", ["Tournament"]),
    ("Sun 4/11", "11:00 AM", [0, 1]), ("Sun 4/11", "1:00 PM", [0]),
    ("Sat 4/17", "9:00 AM", [0, 1, 2]), ("Sat 4/17", "11:00 AM", [0, 1, 2]),
    ("Sat 4/24", "9:00 AM", ["Reserved"]), ("Sat 4/24", "11:00 AM", [0, 1]), ("Sat 4/24", "1:00 PM", [0]),
]

out = os.path.join(os.path.dirname(__file__), "..", "public", "tutorial")
os.makedirs(out, exist_ok=True)
for sport, (facility, heading, units, time_word) in SPORTS.items():
    rows = [["Date", "Start", heading]]
    for date, start, cells in PLAN:
        for c in cells:
            rows.append([date, start, units[c] if isinstance(c, int) else c])
    wb = Workbook()
    ws = wb.active
    ws.title = "April 2027"
    ws.append([f"{facility}: April 2027 {time_word}"])
    ws.merge_cells("A1:C1")
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    for r in rows:
        ws.append(r)
    for c in ws[3]:
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["C"].width = 14
    wb.save(os.path.join(out, f"{sport}-april-2027.xlsx"))
    with open(os.path.join(out, f"{sport}-april-2027.csv"), "w", newline="") as f:
        csv.writer(f).writerows(rows)
print("wrote", len(SPORTS) * 2, "files to", os.path.normpath(out))
